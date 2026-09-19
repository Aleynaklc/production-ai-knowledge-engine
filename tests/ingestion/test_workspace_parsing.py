"""Regression tests for reading order, margin text, and explicit OCR behavior."""

from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.app.config import Settings
from backend.app.documents.validation import DocumentUploadError
from backend.app.workspaces.parsing import extract_in_process


def test_docx_retains_margins_and_groups_related_body_paragraphs() -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header reference ABC-17"
    document.sections[0].footer.paragraphs[0].text = "Footer reference XYZ-29"
    document.add_heading("Eligibility", 1)
    document.add_paragraph("Applicants must be over 18.")
    document.add_paragraph("They must provide proof of age.")
    output = BytesIO()
    document.save(output)
    units = extract_in_process("policy.docx", output.getvalue(), Settings())
    assert any("over 18" in unit.text and "proof of age" in unit.text for unit in units)
    assert any("ABC-17" in unit.text for unit in units)
    assert any("XYZ-29" in unit.text for unit in units)
    assert all(unit.kind == "section" for unit in units)


def test_pdf_layout_orders_text_by_page_position() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 50 200 Td (BOTTOM FACT) Tj ET BT /F1 12 Tf 50 700 Td (TOP HEADING) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    text = extract_in_process("layout.pdf", output.getvalue(), Settings())[0].text
    assert text.index("TOP HEADING") < text.index("BOTTOM FACT")


def test_image_only_page_requires_successful_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    class ImagePage:
        images = [object()]

        def get_contents(self) -> None:
            return None

        def get(self, key: str) -> bool:
            return True

    monkeypatch.setattr(
        "backend.app.workspaces.parsing.PdfReader",
        lambda _: SimpleNamespace(is_encrypted=False, pages=[ImagePage()]),
    )
    monkeypatch.setattr(
        "backend.app.workspaces.parsing._ocr_page", lambda *_: "Scanned answer is 42."
    )
    units = extract_in_process("scan.pdf", b"fixture", Settings())
    assert units[0].number == 1 and units[0].text == "Scanned answer is 42."
    with pytest.raises(DocumentUploadError, match="OCR"):
        extract_in_process("scan.pdf", b"fixture", Settings(upload_ocr_enabled=False))
    monkeypatch.setattr("backend.app.workspaces.parsing._ocr_page", lambda *_: "")
    with pytest.raises(DocumentUploadError, match="OCR"):
        extract_in_process("scan.pdf", b"fixture", Settings())
