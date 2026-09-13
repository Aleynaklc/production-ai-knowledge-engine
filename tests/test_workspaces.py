"""Cross-company security, durable ingestion, version lineage, and incremental work."""

import re
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from threading import Event
from time import monotonic, sleep

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.app.config import Settings
from backend.app.documents.library import DocumentUploadError
from backend.app.llm.generation import GenerationOptions, GenerationResult
from backend.app.workspaces.api import create_app
from backend.app.workspaces.engine import SharedModels, WorkspaceManager
from backend.app.workspaces.parsing import SourceUnit, extract


class Codec:
    def offsets(self, text: str) -> list[tuple[int, int]]:
        return [(match.start(), match.end()) for match in re.finditer(r"\S+", text)]

    def count(self, text: str) -> int:
        return len(self.offsets(text))


class Embedding:
    dimension = 2

    def __init__(self) -> None:
        self.embedded: list[str] = []

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.embedded.extend(texts)
        return [[1.0, 0.0] for _ in texts]

    def embed_query(self, text: str) -> list[float]:
        return [1.0, 0.0]


class Scorer:
    def score(self, query: str, passages: list[str]) -> list[float]:
        return [10.0] * len(passages)


class Generator:
    calls = 0

    def generate(self, prompt: str, system_prompt: str) -> GenerationResult:
        self.calls += 1
        match = re.search(r"Access code is ([A-Z0-9-]+)", prompt)
        text = f"Access code is {match[1]} [S1]." if match else "Insufficient context."
        return GenerationResult(
            text=text,
            input_tokens=30,
            output_tokens=9,
            generation_seconds=0.01,
            tokens_per_second=900,
            options=GenerationOptions(do_sample=False),
        )


@dataclass
class Actor:
    id: str
    workspace: str
    headers: dict[str, str]
    email: str


def register(client: TestClient, email: str) -> Actor:
    response = client.post(
        "/api/v1/auth/register",
        json={
            "email": email,
            "password": "a-secure-test-password",
            "workspace_name": email.split("@")[0],
        },
    )
    assert response.status_code == 200, response.text
    data = response.json()
    workspace = data["workspaces"][0]["id"]
    return Actor(
        data["user"]["id"],
        workspace,
        {
            "Cookie": f"pake_session={response.cookies['pake_session']}",
            "X-Workspace-ID": workspace,
        },
        email,
    )


@pytest.fixture
def api(tmp_path: Path) -> Iterator[tuple[TestClient, WorkspaceManager, Embedding, Generator]]:
    settings = Settings(
        environment="test",
        workspaces_path=tmp_path,
        rag_preload_on_startup=True,
        rag_min_retrieval_score=0,
        chunk_size_tokens=32,
        chunk_overlap_tokens=4,
    )
    embedding, generator = Embedding(), Generator()
    manager = WorkspaceManager(settings, SharedModels(embedding, Scorer(), generator, Codec()))
    with TestClient(create_app(settings, manager)) as client:
        yield client, manager, embedding, generator


type API = tuple[TestClient, WorkspaceManager, Embedding, Generator]
QUESTION = {"question": "What is the access code?"}


def upload(client: TestClient, actor: Actor, text: bytes, filename: str = "guide.txt") -> str:
    response = client.post(
        "/api/v1/documents", params={"filename": filename}, content=text, headers=actor.headers
    )
    assert response.status_code == 202, response.text
    assert response.json()["document"]["status"] == "processing"
    return str(response.json()["document"]["id"])


def wait_document(
    client: TestClient, actor: Actor, document_id: str, status: str = "ready"
) -> dict[str, object]:
    deadline = monotonic() + 5
    while monotonic() < deadline:
        result = client.get("/api/v1/documents", headers=actor.headers)
        assert result.status_code == 200, result.text
        record = next(item for item in result.json()["documents"] if item["id"] == document_id)
        if record["status"] == status:
            return dict(record)
        if record["status"] == "failed" and status != "failed":
            pytest.fail(str(record))
        sleep(0.01)
    pytest.fail("Document processing did not reach expected state")


def test_two_companies_never_share_documents_answers_caches_sources_or_traces(api: API) -> None:
    client, _, embedding, generator = api
    a, b = register(client, "alpha@example.com"), register(client, "beta@example.com")
    a_id = upload(client, a, b"Access code is ALPHA-111. This is private company information.")
    b_id = upload(client, b, b"Access code is BETA-222. This is private company information.")
    wait_document(client, a, a_id)
    wait_document(client, b, b_id)
    a_first = client.post("/api/v1/answers", json=QUESTION, headers=a.headers).json()
    b_first = client.post("/api/v1/answers", json=QUESTION, headers=b.headers).json()
    assert "ALPHA-111" in a_first["result"]["answer"]
    assert "BETA-222" in b_first["result"]["answer"]
    assert not a_first["result"]["cache_hit"] and not b_first["result"]["cache_hit"]
    assert generator.calls == 2 and len(embedding.embedded) == 2
    for actor in (a, b):
        assert client.post("/api/v1/answers", json=QUESTION, headers=actor.headers).json()[
            "result"
        ]["cache_hit"]
        assert client.get("/api/v1/documents", headers=actor.headers).json()["count"] == 1
    assert client.get(f"/api/v1/traces/{a_first['trace_id']}", headers=b.headers).status_code == 404
    assert (
        client.get(f"/api/v1/documents/{a_id}/source?version=1", headers=b.headers).status_code
        == 404
    )
    assert (
        client.get(f"/api/v1/documents/{a_id}/download?version=1", headers=b.headers).status_code
        == 404
    )
    assert client.delete(f"/api/v1/documents/{a_id}", headers=b.headers).status_code == 404
    assert (
        client.put(
            f"/api/v1/documents/{a_id}?filename=hack.txt", content=b"attack", headers=b.headers
        ).status_code
        == 404
    )
    forged = {**b.headers, "X-Workspace-ID": a.workspace}
    assert client.post("/rag/answer", json=QUESTION, headers=forged).status_code == 404
    assert client.get("/api/v1/documents", headers=forged).status_code == 404


def test_membership_roles_revocation_and_session_security(api: API) -> None:
    client, manager, _, _ = api
    owner, reader = register(client, "owner@example.com"), register(client, "reader@example.com")
    assert (
        client.post(
            "/api/v1/workspace/members",
            headers=owner.headers,
            json={"email": reader.email, "role": "reader"},
        ).status_code
        == 200
    )
    reader_headers = {**reader.headers, "X-Workspace-ID": owner.workspace}
    assert client.get("/api/v1/documents", headers=reader_headers).status_code == 200
    assert (
        client.post(
            "/api/v1/documents?filename=x.txt", content=b"forbidden", headers=reader_headers
        ).status_code
        == 403
    )
    assert (
        client.post(
            "/api/v1/workspace/members",
            headers=reader_headers,
            json={"email": owner.email, "role": "editor"},
        ).status_code
        == 403
    )
    assert (
        client.delete(f"/api/v1/workspace/members/{reader.id}", headers=owner.headers).status_code
        == 200
    )
    assert client.get("/api/v1/documents", headers=reader_headers).status_code == 404
    with manager.accounts.connect() as db:
        stored = db.execute("SELECT password FROM accounts WHERE id=?", (owner.id,)).fetchone()[0]
        assert stored != b"a-secure-test-password"
        assert not db.execute(
            "SELECT 1 FROM sessions WHERE token_hash=?", (owner.headers["Cookie"].split("=", 1)[1],)
        ).fetchone()
    assert client.post("/api/v1/auth/logout", headers=owner.headers).status_code == 200
    assert client.get("/api/v1/auth/me", headers=owner.headers).status_code == 401
    response = client.post(
        "/api/v1/auth/login", json={"email": owner.email, "password": "a-secure-test-password"}
    )
    assert response.status_code == 200
    assert "httponly" in response.headers["set-cookie"].lower()
    assert "samesite=lax" in response.headers["set-cookie"].lower()
    with manager.accounts.connect() as db:
        db.execute("UPDATE sessions SET expires=0")
    assert client.get("/api/v1/auth/me").status_code == 401


def test_authentication_csrf_and_rate_limits(api: API) -> None:
    client, _, _, _ = api
    for path in ("/api/v1/documents", "/api/v1/traces", "/api/v1/system"):
        assert client.get(path).status_code == 401
    assert client.post("/api/v1/answers", json=QUESTION).status_code == 401
    assert client.post("/rag/answer", json=QUESTION).status_code == 401
    assert (
        client.post(
            "/api/v1/auth/register",
            headers={"Origin": "https://evil.example"},
            json={
                "email": "x@example.com",
                "password": "long-test-password",
                "workspace_name": "x",
            },
        ).status_code
        == 403
    )
    for _ in range(10):
        assert (
            client.post(
                "/api/v1/auth/login", json={"email": "none@example.com", "password": "wrong"}
            ).status_code
            == 401
        )
    assert (
        client.post(
            "/api/v1/auth/login", json={"email": "none@example.com", "password": "wrong"}
        ).status_code
        == 429
    )


def test_incremental_update_delete_and_versioned_sources(api: API) -> None:
    client, manager, embedding, _ = api
    actor = register(client, "versions@example.com")
    first = upload(client, actor, b"Access code is OLD-100. This document defines the access code.")
    wait_document(client, actor, first)
    second = upload(client, actor, b"Other reference material for employees.")
    wait_document(client, actor, second)
    before = len(embedding.embedded)
    response = client.put(
        f"/api/v1/documents/{first}?filename=updated.txt",
        content=b"Access code is NEW-200. This document defines the access code.",
        headers=actor.headers,
    )
    assert response.status_code == 202
    wait_document(client, actor, first)
    assert len(embedding.embedded) == before + 1
    old = client.get(f"/api/v1/documents/{first}/source?version=1", headers=actor.headers).json()
    new = client.get(f"/api/v1/documents/{first}/source?version=2", headers=actor.headers).json()
    assert "OLD-100" in old["unit"]["text"] and old["filename"] == "guide.txt"
    assert "NEW-200" in new["unit"]["text"] and new["filename"] == "updated.txt"
    assert all(
        "OLD-100" not in chunk.text for chunk in manager.engine(actor.workspace).store.snapshot()[1]
    )
    assert client.delete(f"/api/v1/documents/{first}", headers=actor.headers).status_code == 200
    assert (
        client.get(f"/api/v1/documents/{first}/source?version=1", headers=actor.headers).status_code
        == 404
    )
    assert all(
        chunk.document_id != first for chunk in manager.engine(actor.workspace).store.snapshot()[1]
    )
    assert len(embedding.embedded) == before + 1


def test_processing_does_not_block_answers_and_failed_update_keeps_ready_version(
    api: API, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, _, _, _ = api
    actor = register(client, "background@example.com")
    document_id = upload(
        client, actor, b"Access code is KEEP-100. This is the current working code."
    )
    wait_document(client, actor, document_id)
    started, release = Event(), Event()
    original = extract

    def blocked(filename: str, data: bytes, settings: Settings) -> list[SourceUnit]:
        started.set()
        assert release.wait(5)
        return original(filename, data, settings)

    monkeypatch.setattr("backend.app.workspaces.engine.extract", blocked)
    try:
        response = client.put(
            f"/api/v1/documents/{document_id}?filename=broken.pdf",
            content=b"not a pdf",
            headers=actor.headers,
        )
        assert response.status_code == 202
        assert started.wait(2)
        answer = client.post("/api/v1/answers", json=QUESTION, headers=actor.headers)
        assert answer.status_code == 200 and "KEEP-100" in answer.json()["result"]["answer"]
    finally:
        release.set()
    record = wait_document(client, actor, document_id, "failed")
    assert record["active_version"] == 1
    assert (
        "KEEP-100"
        in client.post("/api/v1/answers", json=QUESTION, headers=actor.headers).json()["result"][
            "answer"
        ]
    )
    assert (
        client.post(f"/api/v1/documents/{document_id}/retry", headers=actor.headers).status_code
        == 202
    )
    wait_document(client, actor, document_id, "failed")


def pdf_bytes() -> bytes:
    writer = PdfWriter()
    for text in ("Company introduction.", "Access code is PDF-222. This code belongs to page two."):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_pages_docx_sections_and_failed_scans(api: API) -> None:
    client, manager, _, _ = api
    actor = register(client, "formats@example.com")
    document_id = upload(client, actor, pdf_bytes(), "guide.pdf")
    wait_document(client, actor, document_id)
    source = client.get(
        f"/api/v1/documents/{document_id}/source?version=1&unit=2", headers=actor.headers
    ).json()
    assert source["unit"]["kind"] == "page" and "PDF-222" in source["unit"]["text"]
    chunks = manager.engine(actor.workspace).store.snapshot()[1]
    assert any(chunk.metadata["source_unit"] == 2 and "PDF-222" in chunk.text for chunk in chunks)
    doc = Document()
    doc.add_heading("Access policy", 0)
    doc.add_paragraph("Access code is DOCX-333.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Department"
    table.cell(0, 1).text = "Support"
    output = BytesIO()
    doc.save(output)
    units = extract("policy.docx", output.getvalue(), manager.settings)
    assert all(unit.kind == "section" for unit in units)
    assert any("Department | Support" in unit.text for unit in units)
    docx_id = upload(client, actor, output.getvalue(), "policy.docx")
    wait_document(client, actor, docx_id)
    blank = PdfWriter()
    blank.add_blank_page(100, 100)
    blank_output = BytesIO()
    blank.write(blank_output)
    with pytest.raises(DocumentUploadError, match="OCR"):
        extract("scan.pdf", blank_output.getvalue(), manager.settings)


def test_queue_and_accounts_survive_restart(tmp_path: Path) -> None:
    settings = Settings(workspaces_path=tmp_path, rag_preload_on_startup=True)
    shared = SharedModels(Embedding(), Scorer(), Generator(), Codec())
    first = WorkspaceManager(settings, shared)
    user = first.accounts.register("restart@example.com", "long-test-password", "Restart")
    workspace = first.accounts.workspaces(user)[0]["id"]
    store = first.store(workspace)
    record, _ = store.enqueue("guide.txt", b"Access code is RESTART-123.")
    assert store.next_job() is not None  # Simulate process interruption after claiming a job.
    second = WorkspaceManager(settings, shared)
    with TestClient(create_app(settings, second)) as client:
        session = client.post(
            "/api/v1/auth/login",
            json={"email": "restart@example.com", "password": "long-test-password"},
        )
        actor = Actor(
            user,
            workspace,
            {
                "Cookie": f"pake_session={session.cookies['pake_session']}",
                "X-Workspace-ID": workspace,
            },
            "restart@example.com",
        )
        wait_document(client, actor, record.id)
        assert second.engine(workspace).store.snapshot()[1]
    third = WorkspaceManager(settings, shared)
    with TestClient(create_app(settings, third)):
        assert third.store(workspace).get(record.id).status == "ready"
        assert third.engine(workspace).store.snapshot()[1]


def test_failed_index_publication_keeps_old_evidence(
    api: API, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, manager, _, _ = api
    actor = register(client, "rollback@example.com")
    document_id = upload(client, actor, b"Access code is OLD-100. This is the current code.")
    wait_document(client, actor, document_id)
    engine = manager.engine(actor.workspace)

    def fail_commit(*args: object, **kwargs: object) -> bool:
        raise RuntimeError("simulated persistence failure")

    monkeypatch.setattr(engine.store, "complete", fail_commit)
    assert (
        client.put(
            f"/api/v1/documents/{document_id}?filename=new.txt",
            content=b"Access code is NEW-999. This is the replacement code.",
            headers=actor.headers,
        ).status_code
        == 202
    )
    record = wait_document(client, actor, document_id, "failed")
    assert record["active_version"] == 1
    answer = client.post("/api/v1/answers", json=QUESTION, headers=actor.headers).json()["result"]
    assert "OLD-100" in answer["answer"] and "NEW-999" not in answer["answer"]


def test_delete_during_processing_cannot_resurrect_document(
    api: API, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, manager, _, _ = api
    actor = register(client, "delete-job@example.com")
    started, release, finished = Event(), Event(), Event()
    original = extract

    def blocked(filename: str, data: bytes, settings: Settings) -> list[SourceUnit]:
        started.set()
        try:
            assert release.wait(5)
            return original(filename, data, settings)
        finally:
            finished.set()

    monkeypatch.setattr("backend.app.workspaces.engine.extract", blocked)
    document_id = upload(client, actor, b"Access code is DELETE-100.")
    try:
        assert started.wait(2)
        assert (
            client.delete(f"/api/v1/documents/{document_id}", headers=actor.headers).status_code
            == 200
        )
    finally:
        release.set()
    assert finished.wait(2)
    engine = manager.engine(actor.workspace)
    with engine.lock:
        assert engine.store.documents() == []
        assert engine.store.snapshot()[1] == []
    assert client.post("/api/v1/answers", json=QUESTION, headers=actor.headers).status_code == 409


def test_only_one_worker_can_own_the_workspace_directory(api: API) -> None:
    _, manager, _, _ = api
    other = WorkspaceManager(manager.settings, manager.shared)
    try:
        with pytest.raises(RuntimeError, match="one worker"):
            other.start()
    finally:
        other.close()
