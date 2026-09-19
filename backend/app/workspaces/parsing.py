"""Bounded PDF/DOCX/text extraction with honest source locations."""

import base64
import json
import shutil
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel
from pypdf import PdfReader

from backend.app.config import Settings
from backend.app.documents.validation import DocumentUploadError, _decode_document


class SourceUnit(BaseModel):
    number: int
    kind: str
    text: str


def _ocr_page(data: bytes, number: int, languages: str) -> str:
    """OCR one image-only page locally, with bounded raster size and subprocess time."""
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        raise DocumentUploadError(
            "ocr_unavailable",
            "This PDF needs OCR. Install Poppler and Tesseract on the server.",
            422,
        )
    with TemporaryDirectory(prefix="pake-ocr-") as directory:
        pdf = Path(directory) / "source.pdf"
        output = Path(directory) / "page"
        pdf.write_bytes(data)
        subprocess.run(
            [
                "pdftoppm",
                "-f",
                str(number),
                "-l",
                str(number),
                "-scale-to",
                "2400",
                "-png",
                "-singlefile",
                str(pdf),
                str(output),
            ],
            check=True,
            capture_output=True,
            timeout=15,
        )
        result = subprocess.run(
            ["tesseract", str(output) + ".png", "stdout", "-l", languages],
            check=True,
            capture_output=True,
            text=True,
            timeout=15,
        )
        return result.stdout.strip()


def extract(filename: str, data: bytes, settings: Settings) -> list[SourceUnit]:
    """Keep untrusted binary parsing out of the API process, with a hard timeout."""
    if Path(filename).suffix.lower() not in {".pdf", ".docx"}:
        return extract_in_process(filename, data, settings)
    payload = {
        "filename": filename,
        "data": base64.b64encode(data).decode(),
        "max_bytes": settings.upload_max_expanded_bytes,
        "max_pages": settings.upload_max_pages,
        "ocr_enabled": settings.upload_ocr_enabled,
        "ocr_languages": settings.upload_ocr_languages,
    }
    try:
        result = subprocess.run(
            [sys.executable, "-m", "backend.app.workspaces.parser_worker"],
            input=json.dumps(payload),
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        if result.returncode not in {0, 1} or not result.stdout:
            raise ValueError("Parser stopped")
        parsed = json.loads(result.stdout)
        if "error" in parsed:
            raise DocumentUploadError(parsed["error"]["code"], parsed["error"]["message"], 422)
        return [SourceUnit.model_validate(item) for item in parsed["units"]]
    except (subprocess.TimeoutExpired, OSError, ValueError) as error:
        if isinstance(error, DocumentUploadError):
            raise
        raise DocumentUploadError(
            "processing_limit", "Document exceeded parser limits or could not be read.", 422
        ) from error


def extract_in_process(filename: str, data: bytes, settings: Settings) -> list[SourceUnit]:
    """Preserve PDF pages; DOCX paragraphs/tables are sections, never invented pages."""
    suffix = Path(filename).suffix.lower()
    units: list[SourceUnit] = []
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF")
            if len(reader.pages) > settings.upload_max_pages:
                raise ValueError("Too many pages")
            for number, page in enumerate(reader.pages, 1):
                contents = page.get_contents()
                if (
                    contents is not None
                    and len(contents.get_data()) > settings.upload_max_expanded_bytes
                ):
                    raise ValueError("Expanded page exceeds limit")
                text = (
                    (page.extract_text(extraction_mode="layout") or "")
                    if contents is not None
                    else ""
                )
                if not text.strip() and page.get("/Resources") and page.images:
                    if not settings.upload_ocr_enabled:
                        raise DocumentUploadError(
                            "ocr_required",
                            "An image-only page needs OCR; enable server OCR before upload.",
                            422,
                        )
                    text = _ocr_page(data, number, settings.upload_ocr_languages)
                    if not text.strip():
                        raise DocumentUploadError(
                            "ocr_failed", "OCR could not read an image-only page.", 422
                        )
                units.append(SourceUnit(number=number, kind="page", text=text))
        elif suffix == ".docx":
            with ZipFile(BytesIO(data)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > settings.upload_max_expanded_bytes
                ):
                    raise ValueError("Expanded DOCX exceeds limit")
                if len(archive.infolist()) > 10_000:
                    raise ValueError("Too many DOCX entries")
            document = Document(BytesIO(data))
            body: list[str] = []
            for block in document.iter_inner_content():
                text = block.text if isinstance(block, Paragraph) else ""
                if (
                    isinstance(block, Paragraph)
                    and block.style
                    and block.style.name.startswith("Heading")
                ):
                    if body:
                        units.append(
                            SourceUnit(
                                number=len(units) + 1, kind="section", text="\n\n".join(body)
                            )
                        )
                        body = []
                    text = "## " + text
                if isinstance(block, Table):
                    text = "\n".join(
                        " | ".join(cell.text for cell in row.cells) for row in block.rows
                    )
                if text.strip():
                    body.append(text)
            if body:
                units.append(
                    SourceUnit(number=len(units) + 1, kind="section", text="\n\n".join(body))
                )
            seen: set[str] = set()
            for section in document.sections:
                for margin in (
                    section.header,
                    section.first_page_header,
                    section.even_page_header,
                    section.footer,
                    section.first_page_footer,
                    section.even_page_footer,
                ):
                    parts = [
                        paragraph.text for paragraph in margin.paragraphs if paragraph.text.strip()
                    ]
                    parts.extend(
                        "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                        for table in margin.tables
                    )
                    text = "\n\n".join(parts).strip()
                    if text and text not in seen:
                        seen.add(text)
                        units.append(SourceUnit(number=len(units) + 1, kind="section", text=text))
        else:
            units = [SourceUnit(number=1, kind="section", text=_decode_document(data))]
        if (
            sum(len(unit.text.encode("utf-8")) for unit in units)
            > settings.upload_max_expanded_bytes
        ):
            raise ValueError("Extracted text exceeds limit")
        if not any(unit.text.strip() for unit in units):
            raise DocumentUploadError(
                "no_extractable_text",
                "No readable text found. Scanned PDFs need OCR before upload.",
                422,
            )
        return units
    except DocumentUploadError:
        raise
    except Exception as error:
        raise DocumentUploadError(
            "invalid_document", "Document is damaged, encrypted, or exceeds processing limits.", 422
        ) from error
